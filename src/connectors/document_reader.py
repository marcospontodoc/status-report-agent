"""Leitura local e independente de documentos de projeto."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


class DocumentReadError(ValueError):
    """Erro claro de localização, formato ou extração de documento."""


@dataclass(frozen=True)
class DocumentContent:
    """Texto extraído de um documento, associado à sua fonte original."""

    source_path: Path
    format: str
    text: str


def _read_markdown(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()


def _read_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    for section in document.sections:
        for area in (section.header, section.footer):
            parts.extend(paragraph.text.strip() for paragraph in area.paragraphs if paragraph.text.strip())
    return "\n".join(parts).strip()


_READERS = {
    ".md": ("Markdown", _read_markdown),
    ".markdown": ("Markdown", _read_markdown),
    ".pdf": ("PDF", _read_pdf),
    ".docx": ("DOCX", _read_docx),
}


def read_document(path: str | Path) -> DocumentContent:
    """Extrai texto de Markdown, PDF ou DOCX sem alterar o documento."""

    source_path = Path(path).expanduser()
    if not source_path.exists():
        raise DocumentReadError(f"Arquivo não encontrado: {source_path}")
    if not source_path.is_file():
        raise DocumentReadError(f"O caminho não é um arquivo: {source_path}")

    reader_spec = _READERS.get(source_path.suffix.lower())
    if reader_spec is None:
        supported = ", ".join(sorted(_READERS))
        raise DocumentReadError(
            f"Formato não suportado em '{source_path}'. Extensões aceitas: {supported}."
        )

    format_name, reader = reader_spec
    try:
        text = reader(source_path)
    except Exception as exc:
        raise DocumentReadError(
            f"Não foi possível ler '{source_path}' como {format_name}: {exc}"
        ) from exc

    if not text.strip():
        raise DocumentReadError(f"O arquivo '{source_path}' não contém texto extraível.")

    return DocumentContent(source_path=source_path.resolve(), format=format_name, text=text)
