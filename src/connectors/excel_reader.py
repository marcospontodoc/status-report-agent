"""Leitura somente leitura do modelo Excel de QA."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any
import unicodedata

from openpyxl import load_workbook
from src.qa_schema import BUG_COLUMNS, BUG_SHEET, SCENARIO_COLUMNS, SCENARIO_SHEET


class ExcelReadError(ValueError):
    """Erro ao abrir ou interpretar a estrutura do workbook de QA."""


@dataclass(frozen=True)
class ExcelRow:
    """Linha não vazia da fonte, com número original e campos por cabeçalho."""

    sheet: str
    row_number: int
    values: dict[str, Any]


@dataclass(frozen=True)
class ExcelQAInput:
    """Linhas lidas das duas abas, sem validação de conteúdo."""

    source_path: Path
    scenarios: tuple[ExcelRow, ...]
    bugs: tuple[ExcelRow, ...]


def _header_key(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFC", str(value))
    return " ".join(text.strip().split()).casefold()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _read_sheet_rows(worksheet: Any, sheet_name: str, expected_columns: tuple[str, ...]) -> tuple[ExcelRow, ...]:
    header_cells = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    positions: dict[str, int] = {}
    duplicates: set[str] = set()

    for index, header in enumerate(header_cells):
        key = _header_key(header)
        if not key:
            continue
        if key in positions:
            duplicates.add(key)
        else:
            positions[key] = index

    if duplicates:
        duplicate_names = sorted(duplicates)
        raise ExcelReadError(f"Aba '{sheet_name}': cabeçalhos duplicados: {', '.join(duplicate_names)}.")

    missing = [column for column in expected_columns if _header_key(column) not in positions]
    if missing:
        raise ExcelReadError(f"Aba '{sheet_name}': cabeçalhos obrigatórios ausentes: {', '.join(missing)}.")

    rows: list[ExcelRow] = []
    for row_number, cells in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if all(_is_blank(value) for value in cells):
            continue
        mapped = {
            column: cells[positions[_header_key(column)]]
            if positions[_header_key(column)] < len(cells)
            else None
            for column in expected_columns
        }
        rows.append(ExcelRow(sheet=sheet_name, row_number=row_number, values=mapped))

    return tuple(rows)


def read_qa_workbook(path: str | Path) -> ExcelQAInput:
    """Lê as abas do modelo sem alterar o arquivo de origem.

    As colunas são associadas pelos cabeçalhos, não pela posição. Linhas
    totalmente vazias são ignoradas; cabeçalhos inválidos interrompem a leitura.
    """

    source_path = Path(path).expanduser()
    if not source_path.is_file():
        raise ExcelReadError(f"Arquivo Excel não encontrado: {source_path}")

    try:
        workbook = load_workbook(source_path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl pode expor vários erros de arquivo/ZIP/XML.
        raise ExcelReadError(f"Não foi possível abrir o arquivo Excel '{source_path}': {exc}") from exc

    try:
        required_sheets = (SCENARIO_SHEET, BUG_SHEET)
        missing_sheets = [name for name in required_sheets if name not in workbook.sheetnames]
        if missing_sheets:
            raise ExcelReadError(f"Abas obrigatórias ausentes: {', '.join(missing_sheets)}.")

        return ExcelQAInput(
            source_path=source_path.resolve(),
            scenarios=_read_sheet_rows(workbook[SCENARIO_SHEET], SCENARIO_SHEET, SCENARIO_COLUMNS),
            bugs=_read_sheet_rows(workbook[BUG_SHEET], BUG_SHEET, BUG_COLUMNS),
        )
    finally:
        workbook.close()
